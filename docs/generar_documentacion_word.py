from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run('NefroHC · v46 · ')
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for item in (begin, instr, separate, text, end): run._r.append(item)


def add_inline(paragraph, text):
    # Preserve common Markdown emphasis and inline code; leave links as readable label + URL.
    pattern = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))')
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos: paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith('`'):
            run = paragraph.add_run(token[1:-1]); run.font.name = 'Consolas'; run.font.size = Pt(9)
        else:
            found = re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)', token)
            paragraph.add_run(f'{found.group(1)} ({found.group(2)})')
        pos = match.end()
    if pos < len(text): paragraph.add_run(text[pos:])


def build(source, target):
    lines = source.read_text(encoding='utf-8').splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2); section.right_margin = Cm(2)
    normal = doc.styles['Normal']; normal.font.name = 'Aptos'; normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(38, 50, 56)
    for style_name, size, color in [('Title', 25, '123B5D'), ('Heading 1', 17, '123B5D'), ('Heading 2', 13, '007C91'), ('Heading 3', 11, '007C91')]:
        style = doc.styles[style_name]; style.font.name = 'Aptos Display'; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color)
    add_page_number(section.footer.paragraphs[0])
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith('```'):
            if not in_code:
                in_code = True; code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(.5)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run('\n'.join(code_lines)); run.font.name = 'Consolas'; run.font.size = Pt(8)
                in_code = False
            i += 1; continue
        if in_code:
            code_lines.append(line); i += 1; continue
        if not line.strip() or line.strip() == '---':
            i += 1; continue
        if line.startswith('|') and i + 1 < len(lines) and re.match(r'^\|?\s*[-:| ]+\|?\s*$', lines[i+1]):
            table_lines=[]
            while i < len(lines) and lines[i].startswith('|'):
                table_lines.append([cell.strip() for cell in lines[i].strip('|').split('|')]); i += 1
            header = table_lines[0]; body = [row for row in table_lines[2:] if row]
            table = doc.add_table(rows=1, cols=len(header)); table.style = 'Light Shading Accent 1'
            for cell, text in zip(table.rows[0].cells, header): cell.text = re.sub(r'\*\*|`', '', text)
            for row in body:
                cells = table.add_row().cells
                for cell, text in zip(cells, row): cell.text = re.sub(r'\*\*|`', '', text)
            doc.add_paragraph('')
            continue
        if line.startswith('# '):
            p = doc.add_paragraph(style='Title'); add_inline(p, line[2:].strip())
        elif line.startswith('## '):
            p = doc.add_heading('', level=1); add_inline(p, line[3:].strip())
        elif line.startswith('#### '):
            p = doc.add_heading('', level=3); add_inline(p, line[5:].strip())
        elif line.startswith('### '):
            p = doc.add_heading('', level=2); add_inline(p, line[4:].strip())
        elif re.match(r'^\s*[-*]\s+', line):
            p = doc.add_paragraph(style='List Bullet'); add_inline(p, re.sub(r'^\s*[-*]\s+', '', line))
        elif re.match(r'^\s*\d+\.\s+', line):
            p = doc.add_paragraph(style='List Number'); add_inline(p, re.sub(r'^\s*\d+\.\s+', '', line))
        elif line.startswith('> '):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(.5); add_inline(p, line[2:]);
        else:
            p = doc.add_paragraph(); add_inline(p, line)
        i += 1
    doc.core_properties.title = source.stem.replace('_', ' ').title()
    doc.core_properties.subject = 'NefroHC v46 · producto 0.1.5 · documento de trabajo'
    doc.save(target)


build(ROOT / 'MANUAL_USUARIO.md', ROOT / 'NefroHC_Manual_Usuario.docx')
build(ROOT / 'DOCUMENTACION_TECNICA.md', ROOT / 'NefroHC_Documentacion_Tecnica.docx')
